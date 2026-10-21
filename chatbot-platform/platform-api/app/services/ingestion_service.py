"""Parse user-selected documents. No URL fetching, directory scans, or executable uploads."""
import io
import json
import csv
import uuid
import zipfile
from pathlib import Path
from ..schemas import KnowledgeDocument

MAX_BYTES = 5_000_000
MAX_CHARS = 200_000
FORMATS = ['.txt', '.md', '.csv', '.json', '.pdf', '.docx', '.xlsx']


def parse_document(name, data):
    if len(data) > MAX_BYTES or not data:
        raise ValueError('Tệp rỗng hoặc vượt giới hạn 5 MB')
    ext = Path(name).suffix.lower()
    if ext not in FORMATS:
        raise ValueError('Định dạng chưa hỗ trợ: ' + ext)
    if ext in ('.docx', '.xlsx'):
        try:
            with zipfile.ZipFile(io.BytesIO(data)) as archive:
                if sum(item.file_size for item in archive.infolist()) > 20_000_000:
                    raise ValueError('Nội dung giải nén vượt 20 MB')
        except zipfile.BadZipFile as exc:
            raise ValueError('Tệp Office không hợp lệ') from exc
    def table(rows):
        rows = list(rows)
        if sum(len(row) for row in rows) > 100000:
            raise ValueError('Bảng vượt 100.000 ô')
        if not rows:
            return ''
        lines = ['| ' + ' | '.join(str('' if v is None else v).replace('|', '\\|').replace('\n', ' ') for v in row) + ' |' for row in rows]
        lines.insert(1, '| ' + ' | '.join('---' for _ in rows[0]) + ' |')
        return '\n'.join(lines)
    try:
        if ext in ('.txt', '.md', '.csv', '.json'):
            text = data.decode('utf-8-sig')
            if ext == '.json':
                text = json.dumps(json.loads(text), ensure_ascii=False, indent=2)
            elif ext == '.csv':
                text = table(csv.reader(io.StringIO(text)))
        elif ext == '.pdf':
            from pypdf import PdfReader
            reader = PdfReader(io.BytesIO(data))
            if reader.is_encrypted:
                raise ValueError('PDF có mật khẩu chưa được hỗ trợ')
            if len(reader.pages) > 200:
                raise ValueError('PDF tối đa 200 trang')
            text = '\n\n'.join(f'## Trang {i+1}\n{page.extract_text() or ""}' for i, page in enumerate(reader.pages))
            if not any((page.extract_text() or '').strip() for page in reader.pages):
                raise ValueError('PDF không có lớp văn bản; cần OCR trước khi nạp')
        elif ext == '.docx':
            from docx import Document
            from docx.text.paragraph import Paragraph
            from docx.table import Table
            doc = Document(io.BytesIO(data))
            parts = []
            for element in doc.element.body:
                if element.tag.endswith('}p'):
                    parts.append(Paragraph(element, doc).text)
                elif element.tag.endswith('}tbl'):
                    parts.append(table([[c.text for c in row.cells] for row in Table(element, doc).rows]))
            text = '\n\n'.join(parts)
        else:
            from openpyxl import load_workbook
            workbook = load_workbook(io.BytesIO(data), read_only=True, data_only=True)
            try:
                text = ''
                for sheet in workbook:
                    rows = []
                    for row in sheet.iter_rows(values_only=True):
                        if len(rows) >= 10000:
                            raise ValueError('Sheet vượt 10.000 dòng')
                        rows.append(row)
                    text += '\n\n## ' + sheet.title + '\n' + table(rows)
                    if len(text) > MAX_CHARS:
                        raise ValueError('Văn bản vượt 200.000 ký tự')
            finally:
                workbook.close()
    except ImportError as exc:
        raise ValueError('Thiếu parser trên backend. Cài lại backend/requirements.txt.') from exc
    except ValueError:
        raise
    except Exception as exc:
        raise ValueError('Không đọc được tệp. Kiểm tra định dạng và nội dung.') from exc
    if not text.strip() or '\x00' in text or len(text) > MAX_CHARS:
        raise ValueError('Văn bản rỗng, có ký tự không hợp lệ hoặc vượt 200.000 ký tự')
    return KnowledgeDocument(id=uuid.uuid4().hex, name=Path(name).name, content=text.strip()).model_dump()
