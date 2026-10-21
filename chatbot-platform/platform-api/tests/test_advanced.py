import unittest
import asyncio
import io
from unittest.mock import patch, AsyncMock
import test_configuration as fixtures
from app.schemas import BotConfig, ChatRequest
from app.services.ingestion_service import parse_document
from app.services.knowledge_service import KnowledgeRetriever
from app.services.policy_service import finalize, match_response
from app.services.configuration_service import catalog


class AdvancedTests(unittest.TestCase):
    setUp = fixtures.ConfigurationTests.setUp
    create = fixtures.ConfigurationTests.create
    publish = fixtures.ConfigurationTests.publish

    def test_ingest_authenticated_and_preserves_office_tables(self):
        self.assertEqual(self.client.post('/api/admin/documents/parse?filename=a.md',content=b'abc').status_code,401)
        res = self.client.post('/api/admin/documents/parse?filename=a.md',headers=self.headers,content='Nội dung'.encode())
        self.assertEqual(res.status_code,200,res.text)
        self.assertEqual(res.json()['document']['content'],'Nội dung')
        self.assertIn('| 0 |',parse_document('a.csv',b'Value\n0')['content'])
        from docx import Document
        doc = Document(); doc.add_paragraph('Before'); table=doc.add_table(rows=2,cols=1)
        table.cell(0,0).text='Score';table.cell(1,0).text='0';doc.add_paragraph('After')
        data=io.BytesIO();doc.save(data)
        text=parse_document('a.docx',data.getvalue())['content']
        self.assertLess(text.index('Before'),text.index('Score'))
        self.assertLess(text.index('Score'),text.index('After'))
        from openpyxl import Workbook
        wb=Workbook();wb.active.append(['Score']);wb.active.append([0]);data=io.BytesIO();wb.save(data)
        self.assertIn('| 0 |',parse_document('a.xlsx',data.getvalue())['content'])
        for name,data in [('a.exe',b'x'),('a.docx',b'bad'),('a.txt',b'\x00'),('a.json',b'{')]:
            with self.assertRaises(ValueError):parse_document(name,data)

    def test_retrieval_modes_context_budget_and_scope(self):
        config=BotConfig(documents=[{'id':'a','name':'a','content':'kiểm định ' * 300},
                                   {'id':'secret','name':'private','content':'SECRET kiểm định','route_id':'private'}],
                         routes=[{'id':'private','label':'Private'}], retrieval={'strategy':'bm25','context_max_chars':1000})
        docs=asyncio.run(KnowledgeRetriever().search('A','kiểm định',config))
        self.assertTrue(docs)
        self.assertLessEqual(sum(len(d['content']) for d in docs),1000)
        self.assertNotIn('secret',[d['id'] for d in docs])
        config.retrieval.strategy='semantic'
        async def fake_embed(texts,model,bot_id):return [[1.,0.] for _ in texts]
        with patch('app.services.knowledge_service.embed',side_effect=fake_embed):
            docs=asyncio.run(KnowledgeRetriever().search('A','kiểm định',config))
        self.assertEqual(docs[0]['score'],1.)

    def test_handlers_are_scoped_and_do_not_call_llm(self):
        bot=self.create(routes=[{'id':'help','label':'Help','handler_id':'h'}],
                        handlers=[{'id':'h','label':'H','kind':'static_response','response':'EXACT'}])
        with patch('app.services.chat_service.generate',new_callable=AsyncMock) as provider:
            out=self.client.post(f"/api/admin/bots/{bot['id']}/preview",headers=self.headers,json={'message':'hi','mode':'help'})
            self.assertEqual(out.json()['response'],'EXACT')
            provider.assert_not_called()
        self.assertEqual(self.client.post('/api/admin/bots',headers=self.headers,json={'routes':[{'id':'x','label':'X','handler_id':'h'}]}).status_code,422)
        other=self.create(routes=[{'id':'read','label':'Read','handler_id':'d'}],handlers=[{'id':'d','label':'D','kind':'document_lookup'}],
                          documents=[{'id':'file','name':'file','content':'kiểm định dữ liệu'}])
        out=self.client.post(f"/api/admin/bots/{other['id']}/preview",headers=self.headers,json={'message':'kiểm định','mode':'read'}).json()
        self.assertEqual(out['reason'],'handler_document_lookup')
        self.assertEqual(out['citations'][0]['id'],'file')

    def test_route_model_and_history_override(self):
        bot=self.create(history_limit=0,routes=[{'id':'r','label':'R','model_override':{'provider':'anthropic','name':'route-model','timeout_seconds':15}}])
        with patch('app.services.chat_service.generate',new_callable=AsyncMock,return_value='ok') as provider:
            self.client.post(f"/api/admin/bots/{bot['id']}/preview",headers=self.headers,json={'message':'hi','mode':'r','history':[{'role':'user','content':'old'}]})
            self.assertEqual(provider.call_args.args[0].model.name,'route-model')
            self.assertEqual(provider.call_args.args[1].history,[])

    def test_tvts_rules_cleaner_and_score_matcher(self):
        pack=catalog()['rule_packs'][0]
        self.assertEqual(len(pack['blocks']),15)
        config=BotConfig(postprocessing={'remove_artifacts':True,'strip_topic_metadata':True}, response_rules=[
            {'id':'score','matcher':'tvts_score','response':'Handoff'}])
        self.assertEqual(match_response(config,'821 điểm thì xét tuyển thế nào').response,'Handoff')
        self.assertIsNone(match_response(config,'điểm cộng chứng chỉ'))
        output=finalize(config,'Topic: internal\nSection: x\nChèn hình.\nNội dung')
        self.assertEqual(output,'Nội dung')
        self.assertEqual(finalize(config,'X '+config.fallback_message+' Y'),config.fallback_message)

    def test_pdf_text_and_scan_rejection(self):
        from pypdf import PdfWriter
        from pypdf.generic import DictionaryObject, NameObject, DecodedStreamObject
        writer=PdfWriter();page=writer.add_blank_page(width=300,height=300)
        font=DictionaryObject({NameObject('/Type'):NameObject('/Font'),NameObject('/Subtype'):NameObject('/Type1'),NameObject('/BaseFont'):NameObject('/Helvetica')})
        page[NameObject('/Resources')]=DictionaryObject({NameObject('/Font'):DictionaryObject({NameObject('/F1'):font})})
        stream=DecodedStreamObject();stream.set_data(b'BT /F1 12 Tf 20 250 Td (Quality assurance) Tj ET')
        page[NameObject('/Contents')]=writer._add_object(stream)
        data=io.BytesIO();writer.write(data)
        self.assertIn('Quality assurance',parse_document('text.pdf',data.getvalue())['content'])
        blank=PdfWriter();blank.add_blank_page(width=300,height=300);data=io.BytesIO();blank.write(data)
        with self.assertRaisesRegex(ValueError,'OCR'):
            parse_document('scan.pdf',data.getvalue())
